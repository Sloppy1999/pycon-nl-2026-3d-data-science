"""Mechanical QA checks on the built deck -> deck/out/qa_checks.json.

The deck is a personal talk in the PyCon NL 2026 style (docs/DESIGN.md): Fira Sans and Fira Mono only, colours
from content.yaml meta.palette only (render palette colours only on the label chip swatches), the master's
footer and slide number on every content slide, and nothing of the former company master anywhere in the
package (check no_teqyard; the speaker's affiliation on slide 1 is the one allowed mention).

Status per check: pass | fail (a defect in the deck) | warn (soft rule, judgement needed) |
pending (blocked on an input another part has not delivered yet: results, renders, snippets).

Run:  uv run python deck/qa/checks.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation

QA = Path(__file__).resolve().parent
DECK = QA.parent
TALK = DECK.parent
OUT = Path(__import__("os").environ.get("DECK_OUT_DIR", DECK / "out")).resolve()

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
DECORATIVE = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
sys.path.insert(0, str(DECK))
from lib.style import ALLOWED_FONTS, FONTS, MIN_PT, PALETTE, RENDER, RENDER_EXTRA  # noqa: E402
from lib.tokens import FILES as RESULT_FILES  # noqa: E402

TEMPLATE = DECK / "template" / "pycon_nl_2026_master.pptx"
GEOMETRY_JSON = DECK / "template" / "pycon_nl_2026_master.geometry.json"
THEME = {"+mj-lt": FONTS["major"], "+mn-lt": FONTS["minor"], "+mj-ea": FONTS["major"], "+mn-ea": FONTS["minor"]}
MONO_FONTS = {FONTS["mono"], FONTS["mono_medium"]}
# fonts allowed in the rendered PDF: the deck's Fira files only (every character on the slides has a Fira glyph,
# check glyph_coverage, so no fallback font may appear)
PDF_FONT_RE = re.compile(r"^(FiraSans|FiraMono)-[A-Za-z]+$")
# the one approved fallback: the abstract's coffee cup in the slide 1 joke (speaker's choice, 2026-10-04); Fira Mono
# has no glyph for it, so an emoji font draws it (NotoColorEmoji here, Segoe UI Emoji on Windows)
EMOJI_ALLOWED = {"\u2615"}
EMOJI_FONT_RE = re.compile(r"^(NotoColorEmoji|SegoeUIEmoji|AppleColorEmoji)$")
# font files that hold each typeface's glyphs (deck/qa/fonts/, the files the QA render and the laptop use):
# (typeface, bold) -> file; italic runs use the italic file where one exists
FONT_DIR = QA / "fonts"
GLYPH_FILES = {
    (FONTS["minor"], False): "FiraSans-Regular.ttf", (FONTS["minor"], True): "FiraSans-Bold.ttf",
    (FONTS["title"], False): "FiraSans-SemiBold.ttf", (FONTS["title"], True): "FiraSans-SemiBold.ttf",
    (FONTS["medium"], False): "FiraSans-Medium.ttf", (FONTS["medium"], True): "FiraSans-Medium.ttf",
    (FONTS["mono"], False): "FiraMono-Regular.ttf", (FONTS["mono"], True): "FiraMono-Bold.ttf",
    (FONTS["mono_medium"], False): "FiraMono-Medium.ttf", (FONTS["mono_medium"], True): "FiraMono-Medium.ttf",
}
ITALIC_FILES = {FONTS["minor"]: "FiraSans-Italic.ttf"}
# sample text of the master's example slides (the deck must not show it); the layouts' own prompt texts are
# added at run time from the template
MASTER_STRINGS = ["Speaker name", "Code, support line and source line go here", "Body text at 24 pt",
                  "Content at 20 pt", "Left column, 24 pt", "Right column, 24 pt", "Optional support line",
                  "Source and licence lines, 12 pt", "What this talk covers", "Hook and running example",
                  "Where a Python loop meets compiled code", "Second level at 20 pt", "These five gotchas",
                  "Fitting points in RAM", "Gotchas and takeaways", "Lorem ipsum"]
SMALL_NAMES = ("Source", "!!src", "!!small", "Small print", "Conditions", "Placeholder note", "tile_ref")
MORPH_OBJECTS = ["!!title", "!!code", "!!support", "!!src", "!!render"]
_KEYS = "|".join(RESULT_FILES)
TOKEN_RE = re.compile(rf"\{{(?:{_KEYS})\.[^}}]*\}}|\[(?:{_KEYS})\.[A-Za-z0-9_.\-]+\]")
# Colours allowed in the package's XML: content.yaml meta.palette; the render palette only on the label chip
# swatches (shapes named "Swatch"), the colours inside the render PNGs are not XML.
SCHEME = set(PALETTE.values())
RENDER_COLOURS = set(RENDER.values()) | {v for v in RENDER_EXTRA.values() if isinstance(v, str)} | \
    {c for v in RENDER_EXTRA.values() if isinstance(v, list) for c in v}

# ---- no_teqyard: nothing of the former company master may remain anywhere in the package --------------------
# Strings (case-insensitive) searched in every XML and .rels part (slides, layouts, master, theme, notes, charts,
# document properties). "teqyard" is allowed only inside the slide 1 presenter line (the speaker's affiliation).
TEQYARD_STRINGS = ["teqyard", "corbel", "calibri", "cambria", "consolas", "madeleine", "carlito", "inconsolata",
                   "driven by data",
                   "powered by ai", "präsentationstitel", "dateiname", "vertraulich", "fußbereich", "titelfolie",
                   "kapiteltrenner", "zitat oder kernaussage", "kontakt und abschluss", "überschrift links",
                   "zwei textspalten", "medien links", "inhaltsverzeichnis", "office theme", "bausteine"]
TEQYARD_RUNS = [">Public<"]                       # the former confidentiality footer, as a whole run text
# case-sensitive: the former master's German date field and language tag
TEQYARD_EXACT = ["Stand:", "de-DE"]
# the former company colours (hex, case-insensitive) in any XML part
TEQYARD_HEX = ["0E2841", "F95759", "4766D2", "F1F1F2", "666A71", "1C1D1F", "8C96A3", "334793", "B6C4F2", "41AE49",
               "FFCE2C", "E20000"]
# SHA-256 of the former master's media parts (logo and background images: ppt/media/image1.jpg, image2.png,
# image3.png, image4.png), hashed from the unpacked copy before the master was deleted; no package part may match
TEQYARD_MEDIA_SHA256 = {
    "d60328d328a2a3de293a82670e0ba1f0b6b5136126d085ccd641fb7657544b76",
    "f849a501470544623259cc376e36e5b6845c3ce514549d49bf21392f42826ae6",
    "be678198d6ded600020036900d0f6e50674296fa3c987856628aefe7662536e4",
    "435b2637936c9d5c5f0b7cb6b53b6ef903ad0145c9dd2aa79abd44c151ae528a",
}
# builder code that must hold no colour value of its own (every colour comes from content.yaml via lib/style.py)
CODE_FILES = [DECK / "build_deck.py"] + sorted((DECK / "lib").glob("*.py"))
COLOUR_LITERAL_RE = re.compile(r"""(["'])#?[0-9A-Fa-f]{6}\1|RGBColor\(\s*0x|RGBColor\(\s*\d""")
# Literal numbers left in appendix/notes text after removing tokens: versions, dates, file and tile ids,
# slide references, code and identifiers. Whatever remains is listed for review (warn).
LITERAL_OK = [r"\{[^}]*\}", r"\bv?\d+(?:\.\d+){2,}\S*", r"\b\d{4}-\d{2}-\d{2}\b", r"\b\w*\d\w*_\w+\b",
              r"\b\w+_\w*\d\w*\b", r"\bslides? \d+(?:(?:,| and| to) \d+)*", r"\b[BC]\d\b", r"`[^`]*`",
              r"\b[A-Za-z]+\d+[A-Za-z0-9]*\b", r"\b\d+[A-Za-z]+\d*[A-Za-z]*\b", r"\([^()]*=\s*[\d.]+[^()]*\)",
              r"\b\w+\(\S*\)", r"\b\d+(?:\.\d+)? ?(?:in|pt|px)\b", r"\bclass(?:es)? \d+(?:(?:,| and) \d+)*",
              r"\b\d+ (?:is|ground|building|water|civil)\b", (r"\b(?:Python|CPython|PyTorch|torch|CUDA|LAS|Open3D|SciPy|"
              r"NumPy|open3d|scipy|numpy|laspy|PyVista|pytorch3d|RTX|Ryzen 7|Lucide|AHN|AHN4|AHN5|AHN6) ?v?[\d.]+\+?\w*"),
              r"\b(?:PF|format) ?\d+\b", r"\bk ?= ?\d+", r"\brow \d+\b", r"\brows \d+ and \d+\b", r"\bsection \d+\b",
              r"\bsections \d+ and \d+\b", r"\bC\d\b", r"\bFLAG \d+(?: and \d+)?\b", r"\bgotcha \d+\b", r"\bE\d\w*\b",
              r"\bstep (?:one|two|five|six|\d)\b", r"\b\d+ (?:to \d+ )?gotchas\b", r"\[[^\]]*\]",
              r"\b(?:one|two|three|four|five|six|seven|eight|nine|ten|twelve) \w+", r"\bC\+\+\d+\b", r"c\+\+\d+",
              r"\bRyzen \d\b", r"\b(?:19|20)\d{2}\b", r"\bCC BY(?:-NC)? \d\.\d\b", r"\bworkers=-1\b", r"\bN by \d\b",
              r"\b\d+ (?:minutes|million)\b", r"\b1[ \u00a0]km\b", r"\b\d+ m crop\b",
              r"-?\d+(?:\.\d+)?(?:, -?\d+(?:\.\d+)?){2,}"]
DASH_RE = re.compile("[–—]")


class Report:
    def __init__(self):
        self.checks = []

    def add(self, cid, status, detail, items=None):
        self.checks.append({"id": cid, "status": status, "detail": detail, "items": items or []})

    def summary(self):
        out = {}
        for c in self.checks:
            out[c["status"]] = out.get(c["status"], 0) + 1
        return out


def package_members(zf):
    """(name, bytes) of every part of the package, and of every part inside an embedded package (the chart
    workbooks ppt/embeddings/*.xlsx), named "<embedding>!/<part>"."""
    import io
    for name in zf.namelist():
        data = zf.read(name)
        yield name, data
        if name.endswith((".xlsx", ".xlsm", ".docx", ".pptx")):
            try:
                inner = zipfile.ZipFile(io.BytesIO(data))
            except zipfile.BadZipFile:
                yield name + "!/", b"not a zip package"
                continue
            for sub in inner.namelist():
                yield f"{name}!/{sub}", inner.read(sub)


def glyph_coverage(slides, labels, pptx_path):
    """Every character of every run on the slides, in the notes and in the chart parts must have a glyph in the
    font file of the run's typeface (bold and italic files where they exist). Title runs inherit Fira Sans
    SemiBold, notes and chart text Fira Sans. Returns (status, detail, items)."""
    try:
        from fontTools.ttLib import TTFont
    except ImportError:
        return "pending", "fontTools not importable", []
    cmaps = {}

    def cmap(fname):
        if fname not in cmaps:
            p = FONT_DIR / fname
            cmaps[fname] = set(TTFont(str(p)).getBestCmap()) if p.exists() else None
        return cmaps[fname]

    if any(cmap(f) is None for f in set(GLYPH_FILES.values())):
        return "pending", f"font files missing in {FONT_DIR} (run deck/qa/setup_fonts.py)", []
    bad = {}

    def check(where, text, face, bold, italic=False):
        fname = ITALIC_FILES.get(face) if italic and face in ITALIC_FILES else GLYPH_FILES.get((face, bool(bold)))
        if fname is None:
            bad.setdefault((where, f"typeface {face!r} has no font file"), 1)
            return
        cm = cmap(fname)
        for ch in text:
            if ch in "\n\r\t\v" or ch in EMOJI_ALLOWED:
                continue
            if ord(ch) not in cm:
                k = (where, f"U+{ord(ch):04X} {ch!r} not in {fname}")
                bad[k] = bad.get(k, 0) + 1

    for sl, lab in zip(slides, labels):
        parts = [("slide", sl._element)]
        if sl.has_notes_slide:
            parts.append(("notes", sl.notes_slide._element))
        for kind, root in parts:
            for sp in root.iter(P + "sp", P + "graphicFrame"):
                ph = sp.find(".//p:nvPr/p:ph", NS)
                in_title = ph is not None and ph.get("type") in ("title", "ctrTitle")
                for r in sp.iter(A + "r", A + "fld"):
                    t = r.find(A + "t")
                    if t is None or not t.text:
                        continue
                    rpr = r.find(A + "rPr")
                    lat = rpr.find(A + "latin") if rpr is not None else None
                    face = lat.get("typeface") if lat is not None else None
                    face = THEME.get(face, face) if face else (FONTS["title"] if in_title else FONTS["minor"])
                    bold = rpr is not None and rpr.get("b") in ("1", "true")
                    italic = rpr is not None and rpr.get("i") in ("1", "true")
                    check(f"{lab} {kind}", t.text, face, bold, italic)
    with zipfile.ZipFile(pptx_path) as zf:
        for name in zf.namelist():
            if re.match(r"ppt/charts/chart\d+\.xml$", name):
                root = etree.fromstring(zf.read(name))
                texts = [t.text or "" for t in root.iter(A + "t")]
                texts += [v.text or "" for v in root.iter("{%s}v" % NS["c"])]
                texts += [n.get("formatCode") or "" for n in root.iter("{%s}numFmt" % NS["c"])]
                for tx in texts:
                    check(name, re.sub(r'\[[^\]]*\]|[";#0.,]', "", tx) if tx.startswith("[") else tx,
                          FONTS["minor"], False)
    items = [f"{w}: {msg} (x{n})" for (w, msg), n in sorted(bad.items())]
    return ("pass" if not items else "fail",
            "every character on the slides, in the notes and in the charts has a glyph in its Fira font file"
            if not items else "characters without a glyph in the run's font (a fallback font would draw them)",
            items)


def words(text: str) -> int:
    return len([w for w in re.split(r"\s+", text.strip()) if re.search(r"\w", w)])


def shape_text(el) -> str:
    paras = []
    for p in el.iter(A + "p"):
        paras.append("".join(t.text or "" for t in p.iter(A + "t")))
    return "\n".join(paras)


def run_font(r, default=FONTS["minor"]) -> str:
    rpr = r.find(A + "rPr")
    lat = rpr.find(A + "latin") if rpr is not None else None
    tf = lat.get("typeface") if lat is not None else None
    return THEME.get(tf, tf) if tf else default


def is_pending(el) -> bool:
    c = el.find(".//" + P + "cNvPr")
    return c is not None and ((c.get("name") or "").startswith("PENDING") or c.get("title") == "PENDING")


def cnv(el):
    for path in ("p:nvSpPr/p:cNvPr", "p:nvPicPr/p:cNvPr", "p:nvGrpSpPr/p:cNvPr", "p:nvGraphicFramePr/p:cNvPr",
                 "p:nvCxnSpPr/p:cNvPr"):
        f = el.find(path, NS)
        if f is not None:
            return f
    return None


def has_alt(el) -> bool:
    c = cnv(el)
    if c is None:
        return False
    if (c.get("descr") or "").strip():
        return True
    return any(e.get("uri") == DECORATIVE for e in c.iter(A + "ext"))


def transition_kind(sld) -> str:
    for ac in sld.findall("mc:AlternateContent", NS):
        ch = ac.find("mc:Choice", NS)
        if ch is not None:
            if ch.find(".//p159:morph", NS) is not None:
                return "morph"
            if ch.find(".//p:fade", NS) is not None:
                return "fade"
    t = sld.find("p:transition", NS)
    if t is not None:
        return "fade" if t.find("p:fade", NS) is not None else "other"
    return "none"


def geom(el):
    off = el.find(".//a:xfrm/a:off", NS)
    ext = el.find(".//a:xfrm/a:ext", NS)
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))


EMU = 914400


def _xfrm(el):
    x = el.find("p:spPr/a:xfrm", NS)
    if x is None:
        x = el.find("p:grpSpPr/a:xfrm", NS)
    return x


def _child_map(grp):
    """Function mapping child EMU coords of a group to slide inches."""
    x = grp.find("p:grpSpPr/a:xfrm", NS)
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    choff, chext = x.find("a:chOff", NS), x.find("a:chExt", NS)
    ox, oy, ex, ey = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    cx, cy = int(choff.get("x")), int(choff.get("y"))
    cex, cey = max(1, int(chext.get("cx"))), max(1, int(chext.get("cy")))

    def f(px, py):
        return ((ox + (px - cx) * ex / cex) / EMU, (oy + (py - cy) * ey / cey) / EMU)
    return f


def _rect_of(el, f):
    x = _xfrm(el)
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    x0, y0 = f(int(off.get("x")), int(off.get("y")))
    x1, y1 = f(int(off.get("x")) + int(ext.get("cx")), int(off.get("y")) + int(ext.get("cy")))
    return (x0, y0, x1 - x0, y1 - y0)


def _segment_of(el, f):
    x = _xfrm(el)
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    X, Y, W, H = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
    fh, fv = x.get("flipH") == "1", x.get("flipV") == "1"
    a = (X + (W if fh else 0), Y + (H if fv else 0))
    b = (X + (0 if fh else W), Y + (0 if fv else H))
    return f(*a), f(*b)


def _pictures(shapes):
    """Every picture shape (free pictures and pictures in placeholders), including those inside groups."""
    from pptx.shapes.picture import Picture
    for sh in shapes:
        if sh.shape_type == 6:          # group
            yield from _pictures(sh.shapes)
        elif isinstance(sh, Picture):   # PlaceholderPicture is a Picture too
            yield sh


def label_chip_check(slides, labels, is_app, chain_labels=(), full_labels=()):
    """Re-measures the label chips in the saved PPTX (independent of the builder's own solver)."""
    sys.path.insert(0, str(DECK))
    from lib.chips import Rect, segment_hits_rect, segments_intersect

    bad, per_slide = [], {}
    for sl, lab, app in zip(slides, labels, is_app):
        if app:
            continue
        tree = sl._element.find(".//p:cSld/p:spTree", NS)
        groups = [el for el in tree if el.tag == P + "grpSp" and cnv(el) is not None
                  and (cnv(el).get("name") or "").startswith(("!!lab_", "Label "))]
        if not groups:
            continue
        ident = lambda px, py: (px / EMU, py / EMU)  # noqa: E731
        frame = None
        for el in tree:
            nm = (cnv(el).get("name") if cnv(el) is not None else "") or ""
            if el.tag == P + "pic" and nm in ("!!render", "Render labels", "Render plane"):
                frame = Rect(*_rect_of(el, ident))
        code = None
        for el in tree:
            nm = (cnv(el).get("name") if cnv(el) is not None else "") or ""
            if nm in ("!!code", "Code") and _xfrm(el) is not None:
                code = Rect(*_rect_of(el, ident))
        chips = {}
        for g in groups:
            f = _child_map(g)
            name = cnv(g).get("name")
            parts = {(cnv(c).get("name") if cnv(c) is not None else ""): c for c in g}
            chip = Rect(*_rect_of(parts["Chip"], f)) if "Chip" in parts else None
            seg = _segment_of(parts["Leader"], f) if "Leader" in parts else None
            ring = Rect(*_rect_of(parts["Anchor"], f)) if "Anchor" in parts else None
            anchor = (ring.x + ring.w / 2, ring.y + ring.h / 2) if ring else None
            if chip is None or seg is None:
                bad.append(f"{lab}: {name}: chip or leader missing")
                continue
            chips[name] = (chip, seg, anchor)
        per_slide[lab] = chips
        names = list(chips)
        for nm in names:
            chip, seg, anchor = chips[nm]
            if frame is not None and not chip.within(frame.inflate(-0.05)):
                bad.append(f"{lab}: {nm} touches or leaves the render frame")
            if code is not None and chip.overlaps(code):
                bad.append(f"{lab}: {nm} overlaps the code column")
            for other in names:
                if other == nm:
                    continue
                oc, os_, oa = chips[other]
                if oa and chip.contains(*oa):
                    bad.append(f"{lab}: {nm} covers the {other} anchor")
                if segment_hits_rect(*seg, oc):
                    bad.append(f"{lab}: {nm} leader runs through {other}")
        for i in range(len(names)):
            for j in range(i + 1, len(names)):
                a, b = chips[names[i]], chips[names[j]]
                if a[0].overlaps(b[0]):
                    bad.append(f"{lab}: {names[i]} and {names[j]} overlap")
                if segments_intersect(*a[1], *b[1]):
                    bad.append(f"{lab}: leaders of {names[i]} and {names[j]} cross")
    # Morph chain: a chip on two consecutive Morph slides with chips must not move (the labels_only slides carry the
    # Ground chip only); between the slides with every chip, every chip must be on both sides (content.yaml)
    chain = [lab for lab in chain_labels if lab in per_slide]
    full = set(full_labels)
    moved = 0
    for a, b in zip(chain, chain[1:]):
        for nm in set(per_slide[a]) | set(per_slide[b]):
            ra, rb = per_slide[a].get(nm), per_slide[b].get(nm)
            if ra is None or rb is None:
                if a in full and b in full:
                    bad.append(f"{a} to {b}: {nm} missing on one side (Morph)")
                continue
            d = max(abs(u - v) for u, v in zip((ra[0].x, ra[0].y, ra[0].w, ra[0].h), (rb[0].x, rb[0].y, rb[0].w, rb[0].h)))
            if d > 0.005:
                moved += 1
                bad.append(f"{a} to {b}: {nm} chip moves {d:.3f} in")
    return bad, f"slides {', '.join(per_slide)}; Morph chain {'-'.join(chain)}"


def main():
    content = yaml.safe_load((DECK / "content.yaml").read_text(encoding="utf-8"))
    meta = content["meta"]
    pptx_path = OUT / meta["file_name"]
    if not pptx_path.exists():
        sys.exit(f"{pptx_path} missing")
    prs = Presentation(str(pptx_path))
    specs = list(content["slides"]) + list(content.get("appendix") or [])
    slides = list(prs.slides)
    labels = [str(s["n"]) for s in specs]
    rep = Report()
    if len(slides) != len(specs):
        rep.add("slide_count", "fail", f"{len(slides)} slides in the deck, {len(specs)} in content.yaml")
    else:
        rep.add("slide_count", "pass", f"{len(slides)} slides ({len(content['slides'])} main + "
                                       f"{len(content.get('appendix') or [])} appendix)")
    is_app = [lab.startswith(("B", "C")) for lab in labels]
    # slide roles from content.yaml (numbers change when slides are inserted): Morph chain, the one on-click build,
    # the filmstrip, the closing slide, the label-chip chain
    main_specs = list(content["slides"])
    morph_labels = [str(s_["n"]) for s_ in main_specs if s_.get("transition") == "morph"]
    build_labels = [str(s_["n"]) for s_ in main_specs if (s_.get("hero") or {}).get("build_on_click")]
    film_label = next((str(s_["n"]) for s_ in main_specs if s_.get("build") == "filmstrip"), None)
    closing_label = next((str(s_["n"]) for s_ in main_specs if s_.get("build") == "closing"), None)
    chip_chain = [str(s_["n"]) for s_ in main_specs if s_.get("transition") == "morph" and s_.get("labels_from")]
    chip_full = [str(s_["n"]) for s_ in main_specs if s_.get("transition") == "morph" and s_.get("labels_from")
                 and not s_.get("labels_only")]
    # the master's layout prompt texts (shown by an empty placeholder) count as master strings too
    master_strings = list(MASTER_STRINGS)
    with zipfile.ZipFile(TEMPLATE) as zt:
        for n in zt.namelist():
            if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n):
                for t in re.findall(r"<a:t>([^<]{4,})</a:t>", zt.read(n).decode("utf-8")):
                    if t != meta["footer"] and t not in master_strings:
                        master_strings.append(t)

    # ---------------------------------------------------------------- per slide text checks
    title_long, title_font, support_long, body_long, fonts_bad, size_bad, alt_missing = [], [], [], [], [], [], []
    dash_hits, token_hits, master_hits, empty_ph, pending_shapes, notes_missing, notes_dash = [], [], [], [], [], [], []
    code_bad, footer_bad, lang_bad, size_soft = [], [], [], []
    for sl, lab, app in zip(slides, labels, is_app):
        sld = sl._element
        tree = sld.find(".//p:cSld/p:spTree", NS)
        body_words = 0
        for el in tree.iter(P + "sp", P + "pic", P + "graphicFrame", P + "grpSp", P + "cxnSp"):
            name = (cnv(el).get("name") if cnv(el) is not None else "") or ""
            ph = el.find(".//p:nvPr/p:ph", NS)
            txt = shape_text(el) if el.tag == P + "sp" else ""
            if is_pending(el):
                pending_shapes.append(f"{lab}: {name}: {txt.splitlines()[0] if txt else ''}")
            # alt text: top-level pictures, charts/tables, groups
            if el.getparent() is tree and el.tag in (P + "pic", P + "graphicFrame", P + "grpSp") and not has_alt(el):
                alt_missing.append(f"{lab}: {name}")
            if el.tag != P + "sp":
                continue
            if ph is not None:
                idx, typ = ph.get("idx"), ph.get("type")
                if typ == "dt":
                    footer_bad.append(f"{lab}: date placeholder present")
                if typ in ("ftr", "sldNum"):
                    continue                     # checked in the footer section below
                if not txt.strip() and el.find(".//p:nvPr/p:ph", NS) is not None:
                    empty_ph.append(f"{lab}: {name} (idx {idx}, type {typ})")
            for m in master_strings:
                if m in txt:
                    master_hits.append(f"{lab}: {name}: {m}")
            if DASH_RE.search(txt):
                dash_hits.append(f"{lab}: {name}: {txt[:60]}")
            for t in TOKEN_RE.findall(txt):
                token_hits.append(f"{lab}: {t}")
            is_title = ph is not None and ph.get("type") in ("title", "ctrTitle")
            small = name.startswith(SMALL_NAMES)
            mono_only = True
            for r in el.iter(A + "r"):
                t = (r.find(A + "t").text or "") if r.find(A + "t") is not None else ""
                if not t.strip():
                    continue
                f = run_font(r)
                rpr = r.find(A + "rPr")
                if rpr is not None and rpr.get("lang") not in (None, "en-GB"):
                    lang_bad.append(f"{lab}: {name}: {rpr.get('lang')}")
                if f not in ALLOWED_FONTS:
                    fonts_bad.append(f"{lab}: {name}: {f}")
                if f not in MONO_FONTS:
                    mono_only = False
                sz = rpr.get("sz") if rpr is not None else None
                # title runs inherit the master's title style (Fira Sans SemiBold, size, colour): no overrides
                if is_title and rpr is not None and (rpr.find(A + "latin") is not None or sz is not None
                                                     or rpr.find(A + "solidFill") is not None):
                    title_font.append(f"{lab}: '{t}' overrides the master title style")
                # hero values (the large runs of a Hero shape; their labels are 20 pt Fira Sans)
                if name.startswith("Hero") and sz is not None and int(sz) >= 2800 and f != FONTS["title"]:
                    title_font.append(f"{lab}: hero '{t}' in {f}, not {FONTS['title']}")
                if sz is None or is_pending(el):
                    continue
                pt = int(sz) / 100
                if f in MONO_FONTS:
                    if pt < MIN_PT["code"]:
                        size_bad.append(f"{lab}: {name}: code {pt} pt")
                elif small:
                    if pt < MIN_PT["small"]:
                        size_bad.append(f"{lab}: {name}: small print {pt} pt")
                elif app:
                    if pt < MIN_PT["appendix"]:
                        size_bad.append(f"{lab}: {name}: appendix {pt} pt")
                elif pt < 18:
                    size_bad.append(f"{lab}: {name}: {pt} pt")
                elif pt < MIN_PT["body"]:
                    size_soft.append(f"{lab}: {name}: {pt} pt (label/chart size, body minimum is 20)")
            if name.startswith(("!!code", "Code")) and not is_pending(el):
                bp = el.find(".//a:bodyPr", NS)
                n_lines = len(list(el.iter(A + "p")))
                if bp is not None and bp.get("wrap") != "none":
                    code_bad.append(f"{lab}: {name}: wrap not off")
                if bp is not None and bp.find("a:noAutofit", NS) is None:
                    code_bad.append(f"{lab}: {name}: autofit not off")
                if n_lines > 12:
                    code_bad.append(f"{lab}: {name}: {n_lines} lines (max 12)")
            if is_title:
                if words(txt) > 8:
                    title_long.append(f"{lab}: {words(txt)} words: {txt!r}")
                continue
            if name in ("Support", "!!support") and words(txt) > 15:
                support_long.append(f"{lab}: {words(txt)} words")
            if not small and not mono_only and not is_pending(el) and not name.startswith(("Code", "!!code")):
                body_words += words(txt)
        if not app and body_words > 20:
            body_long.append(f"{lab}: {body_words} words (excluding title, code, small print)")
        # notes
        nt = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        if not app and (not nt.strip() or "[notes pending]" in nt):
            notes_missing.append(lab)
        if DASH_RE.search(nt):
            notes_dash.append(lab)

    def put(cid, items, detail_ok, detail_bad, bad_status="fail"):
        rep.add(cid, "pass" if not items else bad_status, detail_ok if not items else detail_bad, items)

    put("title_words", title_long, "every title has at most 8 words", "titles over 8 words")
    put("support_words", support_long, "every support line has at most 15 words", "support lines over 15 words")
    put("body_words", body_long, "body text at most about 20 words per main slide",
        "slides over about 20 words of body text (diagram labels count)", "warn")
    put("fonts", fonts_bad, f"only {', '.join(ALLOWED_FONTS)} in slide text", "other fonts in slide text")
    rep.add("glyph_coverage", *glyph_coverage(slides, labels, pptx_path))
    put("min_sizes", size_bad, "body >= 20 pt, code >= 18 pt, small print >= 12 pt, appendix >= 14 pt",
        "text below the minimum size")
    put("label_sizes", size_soft, "no text between 18 and 20 pt", "18 to 20 pt text (allowed for chart and "
        "diagram labels only; check each is a label)", "warn")
    put("alt_text", alt_missing, "alt text (or decorative flag) on every picture, chart, table and group",
        "pictures/charts/tables/groups without alt text")
    put("title_style", title_font, f"title runs inherit the master title style; hero numbers in {FONTS['title']}",
        "title runs with their own font/size/colour, or heroes in another font")
    put("no_dashes", dash_hits, "no en or em dashes in slide text", "en/em dashes in slide text")
    put("no_master_strings", master_hits, "no master placeholder strings in slide text", "master strings found")
    put("no_empty_placeholders", empty_ph, "no empty content placeholders (they would show master prompts)",
        "empty placeholders on slides")
    put("notes_present", notes_missing, "speaker notes on every main slide", "main slides without notes")
    put("notes_no_dashes", notes_dash, "no en/em dashes in speaker notes", "en/em dashes in notes", "warn")
    put("code_boxes", code_bad, "code boxes: wrap off, no autofit, at most 12 lines", "code box issues")
    put("lang_en_gb", lang_bad, "all text runs tagged en-GB", "runs not tagged en-GB")
    rep.add("unresolved_tokens", "pass" if not token_hits else "pending",
            "no unresolved tokens" if not token_hits else f"{len(token_hits)} visible [file.key] placeholders "
            "(results/*.json missing or incomplete)", token_hits)
    rep.add("pending_placeholders", "pass" if not pending_shapes else "pending",
            "no pending placeholder boxes" if not pending_shapes else
            f"{len(pending_shapes)} labelled placeholder boxes (renders, snippets, logo, QR, appendix)", pending_shapes)

    # ---------------------------------------------------------------- transitions, hidden, timing
    tr_bad, hid_bad, timing = [], [], []
    for sl, lab, sp, app in zip(slides, labels, specs, is_app):
        k = transition_kind(sl._element)
        want = "morph" if (not app and lab in morph_labels) else "fade"
        if k != want:
            tr_bad.append(f"{lab}: {k}, expected {want}")
        hidden = sl._element.get("show") == "0"
        if hidden != app:
            hid_bad.append(f"{lab}: hidden={hidden}")
        tm = sl._element.find("p:timing", NS)
        if tm is not None:
            timing.append(lab)
    put("transitions", tr_bad, f"Morph exactly on {morph_labels[0]} to {morph_labels[-1]}, fade on every other slide",
        "wrong transitions")
    put("appendix_hidden", hid_bad, "B and C backup slides hidden, no main slide hidden", "hidden flags wrong")
    rep.add("animations", "pass" if timing == build_labels else "fail",
            f"only slide {', '.join(build_labels)} has a build (on-click hero)" if timing == build_labels
            else f"timing on slides {timing}, expected {build_labels}", timing)
    for hb in build_labels:
        sl = slides[labels.index(hb)]
        tm = sl._element.find("p:timing", NS)
        ok = False
        if tm is not None:
            spids = {t.get("spid") for t in tm.iter(P + "spTgt")}
            hero_ids = {str(sh.shape_id) for sh in sl.shapes if sh.name.startswith("Hero")}
            ok = bool(spids) and spids <= hero_ids
            clicks = len([c for c in tm.iter(P + "cTn") if c.get("nodeType") == "clickEffect"])
            ok = ok and clicks == 1
        rep.add("hero_build", "pass" if ok else "fail", f"slide {hb}: one on-click appear on the hero" if ok
                else f"slide {hb} build does not target the hero only")

    # ---------------------------------------------------------------- dark slides (the master's navy layouts)
    geo = json.loads(GEOMETRY_JSON.read_text(encoding="utf-8"))
    dark_layouts = {v["name"] for v in geo["layouts"].values() if v.get("dark")}
    footer_layouts = {v["name"] for v in geo["layouts"].values() if v.get("footer")}
    dark = [lab for sl, lab, app in zip(slides, labels, is_app) if not app and sl.slide_layout.name in dark_layouts]
    main_labels = [lab for lab, app in zip(labels, is_app) if not app]
    adj = [f"{a} and {b}" for a, b in zip(main_labels, main_labels[1:]) if a in dark and b in dark]
    put("dark_not_adjacent", adj, f"dark slides {dark} never adjacent", "adjacent dark slides")

    # ---------------------------------------------------------------- Morph geometry (the Morph chain)
    morph = [(lab, sl) for lab, sl, app in zip(labels, slides, is_app) if not app and lab in morph_labels]
    geo_bad, names_dup, crop_bad = [], [], []
    ref = {}
    for lab, sl in morph:
        tree = sl._element.find(".//p:cSld/p:spTree", NS)
        seen = {}
        for el in tree:
            c = cnv(el)
            if c is None:
                continue
            nm = c.get("name") or ""
            if nm.startswith("!!"):
                seen[nm] = seen.get(nm, 0) + 1
        names_dup += [f"{lab}: {k} x{v}" for k, v in seen.items() if v > 1]
        for nm in MORPH_OBJECTS:
            els = [el for el in tree if cnv(el) is not None and cnv(el).get("name") == nm]
            if not els:
                geo_bad.append(f"{lab}: {nm} missing")
                continue
            g = geom(els[0])
            if g is None and nm == "!!title":
                g = "inherited"
            if nm not in ref:
                ref[nm] = g
            elif g != ref[nm]:
                geo_bad.append(f"{lab}: {nm} geometry {g} differs from {ref[nm]}")
            if nm == "!!render" and els[0].tag == P + "pic":
                src = els[0].find(".//a:srcRect", NS)
                if src is not None and any(int(src.get(k, "0")) != 0 for k in ("l", "t", "r", "b")):
                    crop_bad.append(f"{lab}: !!render cropped {dict(src.attrib)}")
    put("morph_geometry", geo_bad, f"!!title, !!code, !!support, !!src, !!render identical on {morph_labels[0]} to "
        f"{morph_labels[-1]}",
        "Morph object geometry differs or missing")
    put("morph_unique_names", names_dup, "!! names unique on every Morph slide", "duplicate !! names")
    put("render_no_crop", crop_bad, f"renders on {morph_labels[0]} to {morph_labels[-1]} not cropped", "renders cropped")
    render_kinds = [(lab, [el.tag.split('}')[1] for el in sl._element.iter(P + "pic", P + "sp")
                           if cnv(el) is not None and cnv(el).get("name") == "!!render"]) for lab, sl in morph]
    pics = [lab for lab, k in render_kinds if k == ["pic"]]
    rep.add("morph_renders_present", "pass" if len(pics) == len(morph) else "pending",
            f"{len(pics)} of {len(morph)} Morph slides have the real render picture", [f"{l}: {k}" for l, k in render_kinds])
    s7 = slides[labels.index(film_label)] if film_label in labels else None
    if s7 is not None:
        r7 = [sh for sh in s7.shapes if sh.name == "!!render"]
        st = "pass" if r7 and r7[0].shape_type == 13 else ("pending" if r7 else "fail")
        rep.add("filmstrip_morph_target", st, f"slide {film_label}: first thumbnail named !!render"
                + ("" if st == "pass" else " (not a picture yet)" if r7 else " missing"))

    # ---------------------------------------------------------------- label chips (3, 12, 13, 14)
    chip_bad, chip_info = label_chip_check(slides, labels, is_app, chip_chain, chip_full)
    put("label_chips", chip_bad, "label chips: inside their render frame, no chip/chip or chip/leader overlap, "
        "no crossing leaders, clear of the code column; identical on 12 to 14 where the chips were kept "
        f"({chip_info})", "label chip collisions or Morph mismatch")

    # ---------------------------------------------------------------- footer and slide number
    # every slide on a layout with a footer (all but L1 title and L11 closing) carries the footer text exactly and a
    # slide-number field; no date placeholder anywhere
    for sl, lab in zip(slides, labels):
        tree = sl._element.find(".//p:cSld/p:spTree", NS)
        ftr = [el for el in tree.iter(P + "sp") if el.find(".//p:nvPr/p:ph[@type='ftr']", NS) is not None]
        num = [el for el in tree.iter(P + "sp") if el.find(".//p:nvPr/p:ph[@type='sldNum']", NS) is not None]
        if sl.slide_layout.name in footer_layouts:
            if len(ftr) != 1 or shape_text(ftr[0]).strip() != meta["footer"]:
                footer_bad.append(f"{lab}: footer {[shape_text(f) for f in ftr]!r}, expected {meta['footer']!r}")
            if len(num) != 1 or num[0].find(".//a:fld[@type='slidenum']", NS) is None:
                footer_bad.append(f"{lab}: slide number field missing")
        elif ftr or num:
            footer_bad.append(f"{lab}: footer or slide number on a {sl.slide_layout.name} slide")
    put("footer", footer_bad, f"footer '{meta['footer']}' and slide number on every content slide (not on the "
        "title and closing slides), no date", "footer issues")

    # ---------------------------------------------------------------- layouts and theme: the PyCon NL master
    lay_bad = []
    with zipfile.ZipFile(TEMPLATE) as zt:
        tmpl_layouts = [n for n in zt.namelist() if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n)]
    names = [lay.name for lay in prs.slide_layouts]
    want = [v["name"] for _, v in sorted(geo["layouts"].items(), key=lambda kv: int(kv[0]))]
    if names != want:
        lay_bad.append(f"layouts {names}, master has {want}")
    if len(names) != len(tmpl_layouts):
        lay_bad.append(f"{len(names)} layouts, master has {len(tmpl_layouts)}")
    with zipfile.ZipFile(pptx_path) as zf:
        theme = zf.read("ppt/theme/theme1.xml").decode("utf-8")
    if 'clrScheme name="PyCon NL 2026"' not in theme or 'fontScheme name="PyCon NL 2026"' not in theme:
        lay_bad.append("theme colour/font scheme is not 'PyCon NL 2026'")
    faces = set(re.findall(r'<a:latin typeface="([^"]*)"', theme))
    if faces - set(ALLOWED_FONTS) - {""}:
        lay_bad.append(f"theme fonts {sorted(faces - set(ALLOWED_FONTS))}")
    used = sorted({sl.slide_layout.name for sl in slides})
    put("master", lay_bad, f"deck on the PyCon NL 2026 master ({len(names)} layouts; used: {', '.join(used)}), "
        "theme fonts Fira only", "master/theme issues")

    # ---------------------------------------------------------------- package: metadata, add-in, charts
    pkg_bad = []
    with zipfile.ZipFile(pptx_path) as zf:
        names = zf.namelist()
        core = zf.read("docProps/core.xml").decode("utf-8")
        app_xml = zf.read("docProps/app.xml").decode("utf-8")
        if f"<dc:creator>{meta['author']}</dc:creator>" not in core:
            pkg_bad.append("core.xml author not set to " + meta["author"])
        if f"<dc:title>{meta['title']}</dc:title>" not in core:
            pkg_bad.append("core.xml title is not content.yaml meta.title")
        if f"<dc:language>{meta.get('language', 'en-GB')}</dc:language>" not in core:
            pkg_bad.append("core.xml language is not en-GB")
        if "<Company" in app_xml or "<Manager" in app_xml:
            pkg_bad.append("app.xml has a company or manager field")
        if any("webextension" in n for n in names) or "claude.fileId" in "".join(
                zf.read(n).decode("utf-8", "ignore") for n in names if n.endswith((".xml", ".rels"))):
            pkg_bad.append("add-in part or claude.fileId still present")
        chart_fonts, chart_sizes = set(), []
        for n in names:
            if n.startswith("ppt/charts/chart") and n.endswith(".xml"):
                x = etree.fromstring(zf.read(n))
                for lat in x.iter(A + "latin"):
                    chart_fonts.add(THEME.get(lat.get("typeface"), lat.get("typeface")))
                for el in x.iter(A + "defRPr", A + "rPr"):
                    if el.get("sz") and int(el.get("sz")) < MIN_PT["chart"] * 100:
                        chart_sizes.append(f"{n}: {int(el.get('sz')) / 100} pt")
        if chart_fonts - set(ALLOWED_FONTS):
            pkg_bad.append(f"chart fonts {sorted(chart_fonts - set(ALLOWED_FONTS))}")
        if chart_sizes:
            pkg_bad.append(f"chart text below {MIN_PT['chart']} pt: {chart_sizes[:5]}")
    put("package", pkg_bad, f"core: author {meta['author']}, title, en-GB; no company field, no add-in part; "
        f"chart text Fira Sans >= {MIN_PT['chart']} pt", "package issues")

    # ---------------------------------------------------------------- colours: content.yaml meta.palette only
    # in every XML part (slides, charts, layouts, master, theme, notes); render palette colours only on the
    # label chip swatches
    colour_bad = []
    with zipfile.ZipFile(pptx_path) as zf:
        for name, data in package_members(zf):
            if not (name.startswith("ppt/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(data)
            for el in root.iter(A + "srgbClr", "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}color"):
                if el.tag != A + "srgbClr" and not el.get("rgb"):
                    continue                     # spreadsheet colour by theme index (the deck's theme) or auto
                if el.tag != A + "srgbClr":
                    el = etree.Element("x", val=el.get("rgb")[-6:])
                v = (el.get("val") or "").upper()
                if v in SCHEME:
                    continue
                owner = next((a for a in el.iterancestors() if a.tag in (P + "sp", P + "pic", P + "cxnSp")), None)
                oname = (cnv(owner).get("name") if owner is not None and cnv(owner) is not None else "") or ""
                if v in RENDER_COLOURS and oname == "Swatch":
                    continue
                colour_bad.append(f"{name}: #{v}" + (f" ({oname})" if oname else ""))
    put("colours", sorted(set(colour_bad)), "colours: content.yaml meta.palette only (render palette only on the "
        "label chip swatches)", "colours outside the palette")

    # ---------------------------------------------------------------- no_teqyard: nothing of the former master
    tq = []
    presenter = next((s_.get("presenter", "") for s_ in content["slides"] if str(s_["n"]) == "1"), "")
    allowed = len(re.findall("teqyard", presenter, re.I))
    slide1 = None
    if "1" in labels:
        slide1 = slides[labels.index("1")].part.partname.lstrip("/")
    with zipfile.ZipFile(pptx_path) as zf:
        for name, data in package_members(zf):
            if hashlib.sha256(data).hexdigest() in TEQYARD_MEDIA_SHA256:
                tq.append(f"{name}: image of the former company master")
            if not name.endswith((".xml", ".rels")):
                continue
            text = data.decode("utf-8", "ignore")
            low = text.lower()
            if "!/" in name:
                # inside an embedded package (chart workbook): every font name must be a deck font
                for face in re.findall(r'<(?:\w+:)?name val="([^"]*)"|typeface="([^"]*)"', text):
                    face = face[0] or face[1]
                    if face and face not in ALLOWED_FONTS:
                        tq.append(f"{name}: font {face!r}")
            for w in TEQYARD_STRINGS:
                n_hits = low.count(w)
                if w == "teqyard" and name == slide1:
                    n_hits -= allowed            # the speaker's affiliation in the slide 1 presenter line
                    root = etree.fromstring(data)
                    pres = [el for el in root.iter(P + "sp") if cnv(el) is not None and cnv(el).get("name") == "Presenter"]
                    if allowed and not any("teqyard" in shape_text(el).lower() for el in pres):
                        tq.append(f"{name}: affiliation not in the Presenter shape")
                if n_hits > 0:
                    tq.append(f"{name}: '{w}' x{n_hits}")
            for w in TEQYARD_RUNS:
                if w.lower() in low:
                    tq.append(f"{name}: run text {w}")
            for w in TEQYARD_EXACT:
                if w in text:
                    tq.append(f"{name}: '{w}'")
            for h in TEQYARD_HEX:
                # srgbClr val="RRGGBB", CSS-style #RRGGBB and spreadsheet ARGB "FFRRGGBB", as a whole hex token
                if re.search(rf'(?<![0-9A-Fa-f])(?:[0-9A-Fa-f]{{2}})?{h}(?![0-9A-Fa-f])', text, re.I):
                    tq.append(f"{name}: former company colour #{h}")
    put("no_teqyard", sorted(set(tq)), "no former company string, font, colour, image, layout name or field anywhere "
        "in the package, the embedded chart workbooks included (only deck fonts in them; the slide 1 affiliation "
        f"{presenter!r} excepted)", "former company master traces or foreign fonts")

    # ---------------------------------------------------------------- builder code: no colour values of its own
    lit = []
    for f in CODE_FILES:
        for i, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
            if COLOUR_LITERAL_RE.search(line):
                lit.append(f"{f.relative_to(TALK)}:{i}: {line.strip()[:80]}")
    put("code_colour_literals", lit, "no colour literal in deck/build_deck.py or deck/lib/*.py (every colour comes "
        "from content.yaml through lib/style.py)", "colour literals in builder code")

    # ---------------------------------------------------------------- rendered PDF: embedded fonts
    qr = OUT / "qa_render.json"
    if qr.exists():
        r = json.loads(qr.read_text(encoding="utf-8"))
        pdf_fonts = r.get("pdf_fonts") or []
        bad = [f for f in pdf_fonts if not (PDF_FONT_RE.match(f) or EMOJI_FONT_RE.match(f))]
        rep.add("pdf_fonts", "pass" if pdf_fonts and not bad else "fail",
                f"QA render embeds only Fira Sans, Fira Mono and the emoji fallback: {', '.join(pdf_fonts)}"
                if not bad else f"other fonts in the QA render: {bad}", bad)
    else:
        rep.add("pdf_fonts", "pending", "run deck/qa/render.py first")

    # ---------------------------------------------------------------- research citations resolve to docs/research
    sys.path.insert(0, str(DECK))
    from lib.research_refs import resolve, yaml_texts
    research_root = TALK / "docs" / "research"
    found, missing_refs = resolve(yaml_texts(DECK / "notes.yaml", DECK / "appendix.yaml", DECK / "content.yaml"),
                                  research_root)
    put("research_refs", [f"cited, not in docs/research: {m}" for m in missing_refs],
        f"every 'research:' citation resolves to a file in docs/research ({len(found)} files)",
        "research citations without a file in docs/research")

    # ---------------------------------------------------------------- literal numbers in appendix and notes text
    lit = []
    app = yaml.safe_load((DECK / "appendix.yaml").read_text(encoding="utf-8")) or {}
    notes_y = yaml.safe_load((DECK / "notes.yaml").read_text(encoding="utf-8")) or {}

    def scan(where, text):
        t = str(text)
        for pat in LITERAL_OK:
            t = re.sub(pat, " ", t)
        for m in re.finditer(r"(?<![\w.])\d[\d  .,]*\d*(?: ?(?:%|×|ms|s|µs|ns|MB|GB|B|m|cm|km))?", t):
            lit.append(f"{where}: '{m.group(0).strip()}'")

    for e in app.get("appendix") or []:
        for b in e.get("bullets") or []:
            scan(f"{e['n']} bullet", b.get("text", ""))
        for r in (e.get("table") or {}).get("rows") or []:
            for c in r.get("cells") or []:
                scan(f"{e['n']} table", c)
        scan(f"{e['n']} caption", (e.get("table") or {}).get("caption", ""))
        scan(f"{e['n']} small print", e.get("small_print", ""))
        scan(f"{e['n']} notes", e.get("notes", ""))
    for e in notes_y.get("slides") or []:
        for k, v in e.items():
            if k not in ("n", "seconds", "ends_at", "trace", "checks"):
                for item in v if isinstance(v, list) else [v]:
                    scan(f"notes {e['n']} {k}", item)
    put("literal_numbers", lit, "no literal measured numbers in appendix text or notes "
        "(versions, dates, ids and slide references allowed)", "literal numbers to review (each should be a token, "
        "an id, or a quote)", "warn")

    # ---------------------------------------------------------------- licence: CC BY 4.0 linked, source lines present
    # CC BY 4.0 section 3(a)(1)(C): a URI or hyperlink to the licence. Every "CC BY 4.0" in slide text is a run
    # linked to meta.licence_url; slide 24 also prints the readable URI (meta.licence_line).
    lic_url = meta.get("licence_url", "https://creativecommons.org/licenses/by/4.0/")
    R_ID = "{%s}id" % NS["r"]
    link_bad, src_bad, render_slides = [], [], []
    render_sha = {hashlib.sha1(f.read_bytes()).hexdigest(): f.name for f in (TALK / "renders").glob("*.png")
                  if not f.stem.endswith("_gray")}
    for sl, lab in zip(slides, labels):
        rels = sl.part.rels
        tree = sl._element.find(".//p:cSld/p:spTree", NS)
        for el in tree.iter(P + "sp", P + "graphicFrame"):
            txt = shape_text(el)
            want = txt.count("CC BY 4.0")
            if not want:
                continue
            linked = 0
            for r in el.iter(A + "r"):
                t = r.find(A + "t")
                h = r.find("a:rPr/a:hlinkClick", NS)
                if t is None or h is None or (t.text or "") != "CC BY 4.0":
                    continue
                rel = rels.get(h.get(R_ID)) if h.get(R_ID) in rels else None
                if rel is not None and rel.is_external and rel.target_ref == lic_url:
                    linked += 1
            if linked < want:
                name = (cnv(el).get("name") if cnv(el) is not None else "") or ""
                link_bad.append(f"{lab}: {name}: {want - linked} of {want} 'CC BY 4.0' not linked to {lic_url}")
        # every slide that shows a render from renders/ carries the AHN source line (Source / !!src)
        shown = sorted({render_sha[sh.image.sha1] for sh in _pictures(sl.shapes) if sh.image.sha1 in render_sha})
        if shown:
            render_slides.append(lab)
            srcs = [sh for sh in sl.shapes if sh.name in ("Source", "!!src") and sh.has_text_frame
                    and meta["source_ahn"] in sh.text_frame.text]
            if not srcs:
                src_bad.append(f"{lab}: shows {', '.join(shown)} without a Source / !!src line")
    if closing_label in labels and meta.get("licence_line"):
        s24 = slides[labels.index(closing_label)]
        if not any(sh.has_text_frame and meta["licence_line"] in sh.text_frame.text for sh in s24.shapes):
            link_bad.append(f"{closing_label}: readable licence line {meta['licence_line']!r} missing")
    put("licence_links", link_bad, f"every 'CC BY 4.0' in slide text links to {lic_url}; slide {closing_label} prints "
        "the URI",
        "licence links missing")
    put("source_lines", src_bad, "every slide with a picture from renders/ has the AHN Source / !!src line "
        f"(render pictures on slides {', '.join(render_slides) or 'none'})", "render slides without a source line")

    # ---------------------------------------------------------------- speaker placeholders in visible text
    # [ ... ] with a letter inside, and the QR hint: open items for the speaker. Code boxes are skipped (their
    # brackets are Python indexing); Fira Mono runs elsewhere count, because the gotcha table's placeholder
    # tokens are set in the code font.
    ph_re = re.compile(r"\[[^\]\n]*[A-Za-z][^\]\n]*\]")
    code_names = ("Code", "!!code", "Struck code")
    spk = []
    for sl, lab in zip(slides, labels):
        tree = sl._element.find(".//p:cSld/p:spTree", NS)
        for el in tree.iter(P + "sp", P + "graphicFrame"):
            name = (cnv(el).get("name") if cnv(el) is not None else "") or ""
            if name.startswith(code_names):
                continue
            for para in el.iter(A + "p"):
                vis = "".join((r.find(A + "t").text or "") for r in para.iter(A + "r") if r.find(A + "t") is not None)
                for m in ph_re.findall(vis):
                    spk.append(f"{lab}: {name}: {m}")
                if "Set qr.url" in vis:
                    spk.append(f"{lab}: {name}: Set qr.url")
            if name.startswith("Placeholder note"):
                spk.append(f"{lab}: {name}: {shape_text(el)[:60]}")
    rep.add("speaker_placeholders", "pass" if not spk else "pending",
            "no speaker placeholders in visible text" if not spk else
            f"{len(spk)} speaker placeholders still visible (docs/OPEN_ITEMS.md)", spk)

    # ---------------------------------------------------------------- build report pass-through
    br = OUT / "build_report.json"
    if br.exists():
        b = json.loads(br.read_text(encoding="utf-8"))
        v = b.get("validate", {})
        rep.add("validate_py", "pass" if v.get("returncode") == 0 and "PASSED" in (v.get("output") or "")
                else "fail", (v.get("output") or str(v))[-300:])
        if b.get("warnings"):
            rep.add("build_warnings", "warn", f"{len(b['warnings'])} build warnings (see build_report.json)",
                    [f"{w['slide']}: {w['msg']}" for w in b["warnings"]])
    out = {"deck": str(pptx_path), "summary": rep.summary(), "checks": rep.checks}
    (OUT / "qa_checks.json").write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(out["summary"]))
    for c in rep.checks:
        if c["status"] != "pass":
            print(f"[{c['status']}] {c['id']}: {c['detail']}")
            for it in c["items"][:8]:
                print("      ", it)
            if len(c["items"]) > 8:
                print(f"       ... {len(c['items']) - 8} more")


if __name__ == "__main__":
    main()
