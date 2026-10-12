"""Build the PyCon NL 2026 style slide master for the talk "3D Data Science with Python".

Writes (deterministic, re-runnable):
  deck/template/pycon_nl_2026_master.pptx            master, 11 layouts, one example slide per layout, notes master
  deck/template/pycon_nl_2026_master.geometry.json   every layout placeholder and static shape (EMU and inches)

The package is written part by part (no python-pptx default parts, so no foreign theme, font or metadata).
python-pptx re-opens the result at the end as a check. Palette, fonts and layout rules: docs/DESIGN.md.

Run:  cd talk && uv run python deck/template/make_master.py
"""

from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree

HERE = Path(__file__).resolve().parent
OUT = HERE / "pycon_nl_2026_master.pptx"
GEOM_OUT = HERE / "pycon_nl_2026_master.geometry.json"

LANG = "en-GB"
THEME_NAME = "PyCon NL 2026"
FOOTER = "3D Data Science with Python | PyCon NL 2026"

# ------------------------------------------------------------------ palette (pycon-nl.org, fetched 2026-10-03)
NAVY = "1C325B"    # Tailwind "north"
BLUE = "2B6797"    # "north2"
SAGE = "B1CAC5"    # "north3"
WHITE = "FFFFFF"
INK = "263238"     # site text neutral
SLATE = "455A64"   # site neutral
GREY = "9CA3AF"    # site UI grey (placeholder text, borders)
LIGHT = "E0E0E0"   # site light neutral (theme lt2: tiles, rules)

# theme slot -> hex; the clrMap below maps bg1=lt1, tx1=dk1, bg2=lt2, tx2=dk2
CLR_SCHEME = [
    ("dk1", INK), ("lt1", WHITE), ("dk2", NAVY), ("lt2", LIGHT),
    ("accent1", BLUE), ("accent2", SAGE), ("accent3", NAVY), ("accent4", SLATE),
    ("accent5", GREY), ("accent6", LIGHT), ("hlink", BLUE), ("folHlink", BLUE),
]

SANS = "Fira Sans"
SANS_TITLE = "Fira Sans SemiBold"   # weight 600 is its own family name in the static TTFs (name ID 1)
MONO = "Fira Mono"
PANOSE = {SANS: "020B0503050000020004", SANS_TITLE: "020B0603050000020004", MONO: "020B0509050000020004"}
PITCH = {SANS: "34", SANS_TITLE: "34", MONO: "49"}   # variable pitch / fixed pitch, swiss / modern family

# ------------------------------------------------------------------ geometry (EMU)
EMU_PER_IN = 914400
SLIDE_W, SLIDE_H = 12192000, 6858000            # 13.333 x 7.5 in


def E(inches: float) -> int:
    return round(inches * EMU_PER_IN)


M = E(0.6)                         # outer margin
RIGHT = SLIDE_W - M                # 12.733 in
CW = RIGHT - M                     # content width 12.133 in
GAP = E(0.5)                       # column gap
TITLE_Y = E(0.80)                  # title top on every content layout
TITLE_H1 = E(0.60)                 # one line at 32 pt (32 pt x 1.2 = 0.533 in)
TITLE_H2 = E(1.10)                 # two lines at 32 pt
BAR = (E(0.55), E(0.07))           # small accent bar above the title
BAR_Y = E(0.60)
BODY_Y = E(1.70)                   # body top below a one-line title
BODY_Y2 = E(2.15)                  # body top below a two-line title (L7, L8)
FOOT_Y, FOOT_H = E(6.95), E(0.25)  # footer and slide number: 12 pt, bottom edge 7.20 in
NUM_W = E(1.0)

# media frame: aspect exactly 1588:1792 = 397:448, top on the margin, as tall as the content area allows
_K = int((E(6.20)) // 448)
MEDIA_H = 448 * _K                 # 5 668 992 EMU = 6.1997 in
MEDIA_W = 397 * _K                 # 5 023 638 EMU = 5.4939 in
MEDIA_Y = M
BOTTOM = MEDIA_Y + MEDIA_H         # content bottom 6.7997 in (all content layouts)
COL_W = RIGHT - M - MEDIA_W - GAP  # text column next to the media, 6.1395 in
assert MEDIA_W * 1792 == MEDIA_H * 1588

HALF_W = (CW - GAP) // 2           # L5 columns, 5.8167 in
L10_TITLE_W = E(3.90)              # about 16 to 19 characters per line at 32 pt Fira Sans SemiBold
L10_BODY_X = M + L10_TITLE_W + GAP

DARK_BOTTOM = SLIDE_H - M          # 6.9 in: L1 and L11 have no footer
ZONE_W, ZONE_H = E(2.90), E(1.30)  # logo zone
QR_TILE = E(2.90)
QR_PAD = E(0.20)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
XMLNS = f'xmlns:a="{NS["a"]}" xmlns:r="{NS["r"]}" xmlns:p="{NS["p"]}"'
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
DECORATIVE = ('<a:extLst><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
              '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/>'
              '</a:ext></a:extLst>')
SLIDENUM_FLD = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"

# ------------------------------------------------------------------ text property helpers


def fill(colour: str | None) -> str:
    if colour is None:
        return ""
    if colour.startswith("#"):
        return f'<a:solidFill><a:srgbClr val="{colour[1:]}"/></a:solidFill>'
    return f'<a:solidFill><a:schemeClr val="{colour}"/></a:solidFill>'


def font_xml(face: str | None) -> str:
    if face is None:
        return ""
    if face.startswith("+"):
        return f'<a:latin typeface="{face}"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/>'
    return f'<a:latin typeface="{face}" panose="{PANOSE[face]}" pitchFamily="{PITCH[face]}" charset="0"/>'


def rpr(tag="a:rPr", *, sz=None, b=None, colour=None, face=None, kern=None, dirty=False) -> str:
    attrs = f'lang="{LANG}"'
    if sz is not None:
        attrs += f' sz="{int(sz * 100)}"'
    if b is not None:
        attrs += f' b="{1 if b else 0}"'
    if kern is not None:
        attrs += f' kern="{kern}"'
    if dirty:
        attrs += ' dirty="0"'
    inner = fill(colour) + font_xml(face)
    return f"<{tag} {attrs}>{inner}</{tag}>" if inner else f"<{tag} {attrs}/>"


def ppr(lvl: int, *, marL=0, indent=0, algn="l", ln=100, bef=0, aft=0, bullet=None, rp="") -> str:
    """bullet: None (buNone), ("char", "•", colour) or ("num", "arabicPeriod", colour)."""
    if bullet is None:
        bu = "<a:buNone/>"
    elif bullet[0] == "char":
        bu = f'<a:buClr><a:schemeClr val="{bullet[2]}"/></a:buClr><a:buFont typeface="{SANS}"/><a:buChar char="{bullet[1]}"/>'
    else:
        bu = f'<a:buClr><a:schemeClr val="{bullet[2]}"/></a:buClr><a:buFont typeface="+mj-lt"/><a:buAutoNum type="{bullet[1]}"/>'
    return (f'<a:lvl{lvl}pPr marL="{marL}" indent="{indent}" algn="{algn}" defTabSz="457200" rtl="0" '
            f'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
            f'<a:lnSpc><a:spcPct val="{ln * 1000}"/></a:lnSpc><a:spcBef><a:spcPts val="{bef * 100}"/></a:spcBef>'
            f'<a:spcAft><a:spcPts val="{aft * 100}"/></a:spcAft>{bu}{rp}</a:lvl{lvl}pPr>')


def body_pr(anchor="t", wrap=True) -> str:
    return (f'<a:bodyPr vert="horz" wrap="{"square" if wrap else "none"}" lIns="0" tIns="0" rIns="0" bIns="0" '
            f'rtlCol="0" anchor="{anchor}"><a:noAutofit/></a:bodyPr>')


def para(text: str, *, lvl: int = 0, rp: str | None = None) -> str:
    rp = rp if rp is not None else rpr(dirty=True)
    lvl_attr = f'<a:pPr lvl="{lvl}"/>' if lvl else ""
    t = escape(text)
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return f'<a:p>{lvl_attr}<a:r>{rp}<a:t{space}>{t}</a:t></a:r><a:endParaRPr lang="{LANG}" dirty="0"/></a:p>'


def xfrm(x, y, w, h) -> str:
    return f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'


SPTREE_HEAD = ('<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
               '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/>'
               '<a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')


def ph_attr(ph_type: str | None, idx: int | None, sz: str | None = None) -> str:
    a = ""
    if ph_type:
        a += f' type="{ph_type}"'
    if sz:
        a += f' sz="{sz}"'
    if idx is not None:
        a += f' idx="{idx}"'
    return f"<p:ph{a}/>"


# ------------------------------------------------------------------ layout specs
# Placeholder dict keys: role, name, type (None = obj), idx, box (x, y, w, h in EMU), anchor, lvl1 (defRPr kwargs),
# algn, prompt, sample (example-slide paragraphs), plus "fill" for picture placeholders.
# Static dict keys: kind ("bar" | "tile"), name, box, colour, radius.


def bar(x, y, colour):
    return {"kind": "bar", "name": "Accent bar", "box": (x, y, *BAR), "colour": colour}


def title_ph(x, w, h, *, sample, **kw):
    d = {"role": "title", "name": "Title", "type": "title", "idx": None, "box": (x, TITLE_Y, w, h), "anchor": "t",
         "prompt": "Title in sentence case, max 8 words", "sample": sample}
    d.update(kw)
    return d


def footer_phs(dark: bool):
    colour = "accent2" if dark else "accent4"      # sage on navy, slate on white
    return [
        {"role": "footer", "name": "Footer", "type": "ftr", "idx": 11, "sz": "quarter",
         "box": (M, FOOT_Y, E(8.0), FOOT_H), "anchor": "b", "lvl1": {"sz": 12, "colour": colour},
         "prompt": FOOTER, "sample": [FOOTER]},
        {"role": "slide_number", "name": "Slide number", "type": "sldNum", "idx": 12, "sz": "quarter",
         "box": (RIGHT - NUM_W, FOOT_Y, NUM_W, FOOT_H), "anchor": "b", "algn": "r",
         "lvl1": {"sz": 12, "colour": colour}, "prompt": None, "sample": None},
    ]


def media_ph(x, y, w, h, prompt="Render, 1588 × 1792 px"):
    return {"role": "media", "name": "Media", "type": "pic", "idx": 18, "box": (x, y, w, h), "fill": None,
            "prompt": prompt, "sample": None}


LAYOUTS = [
    {   # 1
        "name": "Title", "type": "title", "dark": True, "footer": False,
        "static": [bar(M, E(1.75), "accent2")],
        "ph": [
            {"role": "title", "name": "Title", "type": "ctrTitle", "idx": None, "box": (M, E(1.95), CW, E(1.00)),
             "anchor": "t", "lvl1": {"sz": 54, "colour": "bg1"}, "prompt": "Talk title",
             "sample": ["3D Data Science with Python"]},
            {"role": "subtitle", "name": "Subtitle", "type": "subTitle", "idx": 1, "box": (M, E(3.20), E(10.50), E(0.95)),
             "anchor": "t", "lvl1": {"sz": 24, "colour": "accent2", "face": MONO},
             "prompt": "Subtitle or code line, two lines of Fira Mono 24 pt",
             "sample": ["for point in point_cloud:", "    # 12 million times. see you tomorrow"]},
            {"role": "presenter", "name": "Presenter", "type": "body", "idx": 13,
             "box": (M, DARK_BOTTOM - E(1.20), RIGHT - ZONE_W - E(0.30) - M, E(1.20)), "anchor": "b",
             "lvl1": {"sz": 24, "colour": "bg1"}, "prompt": "Speaker name · role · repo URL",
             "sample": ["Speaker name · role · [repo URL]"]},
            {"role": "logo", "name": "Logo zone", "type": "pic", "idx": 16,
             "box": (RIGHT - ZONE_W, DARK_BOTTOM - ZONE_H, ZONE_W, ZONE_H), "fill": None,
             "prompt": "PyCon NL logo (file and permission pending)", "sample": None},
        ],
    },
    {   # 2
        "name": "Agenda", "type": None, "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, CW, TITLE_H1, sample=["What this talk covers"]),
            {"role": "body", "name": "Agenda list", "type": None, "idx": 1, "box": (M, BODY_Y, CW, BOTTOM - BODY_Y),
             "anchor": "t", "list": "num", "prompt": "Agenda item",
             "sample": ["Hook and running example", "Landscape", "Pipeline", "Performance cliff",
                        "Fitting points in RAM", "Gotchas and takeaways"]},
        ],
    },
    {   # 3
        "name": "Section", "type": "secHead", "dark": True, "footer": True,
        "static": [bar(M, E(2.60), "accent2")],
        "ph": [
            {"role": "title", "name": "Title", "type": "title", "idx": None, "box": (M, E(2.80), CW, E(1.00)),
             "anchor": "t", "lvl1": {"sz": 44, "colour": "bg1"}, "prompt": "Section title",
             "sample": ["Performance cliff"]},
            {"role": "body", "name": "Section line", "type": "body", "idx": 1, "box": (M, E(3.95), CW, E(0.90)),
             "anchor": "t", "lvl1": {"sz": 24, "colour": "accent2"}, "prompt": "What this section answers",
             "sample": ["Where a Python loop meets compiled code"]},
        ],
    },
    {   # 4
        "name": "Title and content", "type": "obj", "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, CW, TITLE_H1, sample=["No single library does the whole job"]),
            {"role": "body", "name": "Content", "type": None, "idx": 1, "box": (M, BODY_Y, CW, BOTTOM - BODY_Y),
             "anchor": "t", "prompt": "Content: one visual, or body text at 24 pt",
             "sample": ["Body text at 24 pt in ink on white. One idea per slide.",
                        ("Second level at 20 pt with a blue bullet", 1)]},
        ],
    },
    {   # 5
        "name": "Two columns", "type": "twoObj", "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, CW, TITLE_H1, sample=["laspy loads, Open3D computes, PyVista draws"]),
            {"role": "left", "name": "Left column", "type": None, "idx": 1, "box": (M, BODY_Y, HALF_W, BOTTOM - BODY_Y),
             "anchor": "t", "prompt": "Left column", "sample": ["Left column, 24 pt"]},
            {"role": "right", "name": "Right column", "type": None, "idx": 2,
             "box": (M + HALF_W + GAP, BODY_Y, HALF_W, BOTTOM - BODY_Y), "anchor": "t", "prompt": "Right column",
             "sample": ["Right column, 24 pt"]},
        ],
    },
    {   # 6
        "name": "Statement", "type": None, "dark": True, "footer": True,
        "static": [bar(M, E(2.05), "accent2")],
        "ph": [
            {"role": "title", "name": "Statement", "type": "title", "idx": None, "box": (M, E(2.25), E(11.00), E(1.70)),
             "anchor": "t", "lvl1": {"sz": 48, "colour": "bg1"}, "prompt": "One statement, max 8 words",
             "sample": ["Loop over tiles, never over points"]},
            {"role": "body", "name": "Support", "type": "body", "idx": 1, "box": (M, E(4.15), E(11.00), E(1.40)),
             "anchor": "t", "lvl1": {"sz": 28, "colour": "accent2"}, "prompt": "Optional support line or code",
             "sample": ["Optional support line, 28 pt sage"]},
        ],
    },
    {   # 7
        "name": "Text left, media right", "type": None, "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, COL_W, TITLE_H2, sample=["1 Load:", "laspy hands you NumPy"]),
            {"role": "text", "name": "Text column", "type": None, "idx": 14, "box": (M, BODY_Y2, COL_W, BOTTOM - BODY_Y2),
             "anchor": "t", "lvl1": {"sz": 20}, "prompt": "Code (18 pt) or text (20 pt), support line, source line",
             "sample": ["Code, support line and source line go here."]},
            media_ph(RIGHT - MEDIA_W, MEDIA_Y, MEDIA_W, MEDIA_H),
        ],
    },
    {   # 8
        "name": "Media left, text right", "type": None, "dark": False, "footer": True,
        "static": [bar(M + MEDIA_W + GAP, BAR_Y, "accent1")],
        "ph": [
            title_ph(M + MEDIA_W + GAP, COL_W, TITLE_H2, sample=["A point cloud is just", "an N×3 array"]),
            {"role": "text", "name": "Text column", "type": None, "idx": 14,
             "box": (M + MEDIA_W + GAP, BODY_Y2, COL_W, BOTTOM - BODY_Y2), "anchor": "t", "lvl1": {"sz": 20},
             "prompt": "Code (18 pt) or text (20 pt), support line, source line",
             "sample": ["Code, support line and source line go here."]},
            media_ph(M, MEDIA_Y, MEDIA_W, MEDIA_H),
        ],
    },
    {   # 9
        "name": "Large media", "type": None, "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, CW, TITLE_H1, sample=["Six steps turn a point cloud into labels"]),
            media_ph(M, BODY_Y, CW, BOTTOM - BODY_Y, prompt="Picture, chart or diagram"),
        ],
    },
    {   # 10
        "name": "Title left, content right", "type": None, "dark": False, "footer": True,
        "static": [bar(M, BAR_Y, "accent1")],
        "ph": [
            title_ph(M, L10_TITLE_W, E(2.60), sample=["These five gotchas cost me weeks"]),
            {"role": "body", "name": "Content", "type": None, "idx": 1,
             "box": (L10_BODY_X, TITLE_Y, RIGHT - L10_BODY_X, BOTTOM - TITLE_Y), "anchor": "t", "lvl1": {"sz": 20},
             "prompt": "Content: table, list or visual, body 20 pt",
             "sample": ["Content at 20 pt: a table, a list or one visual."]},
        ],
    },
    {   # 11
        "name": "Closing", "type": None, "dark": True, "footer": False,
        "static": [bar(M, BAR_Y, "accent2"),
                   {"kind": "tile", "name": "QR tile", "box": (RIGHT - QR_TILE, E(3.10), QR_TILE, QR_TILE),
                    "colour": "bg1", "radius": 6000}],
        "ph": [
            {"role": "title", "name": "Title", "type": "title", "idx": None,
             "box": (M, TITLE_Y, RIGHT - ZONE_W - E(0.30) - M, E(1.10)), "anchor": "t",
             "lvl1": {"sz": 60, "colour": "bg1"}, "prompt": "Closing title", "sample": ["Copy the pipeline"]},
            {"role": "contact", "name": "Contact", "type": "body", "idx": 15,
             "box": (M, E(2.30), RIGHT - QR_TILE - E(0.30) - M, E(3.70)), "anchor": "b",
             "lvl1": {"sz": 24, "colour": "bg1"}, "aft": 4, "prompt": "Repo URL, name, contact",
             "sample": ["[repo URL]", "Speaker name", "[email]"]},
            {"role": "small_print", "name": "Small print", "type": "body", "idx": 19,
             "box": (M, E(6.20), CW, DARK_BOTTOM - E(6.20)), "anchor": "b", "lvl1": {"sz": 12, "colour": "accent2"},
             "aft": 0, "prompt": "Source and licence lines, 12 pt",
             "sample": ["Source and licence lines, 12 pt sage on navy."]},
            {"role": "logo", "name": "Logo zone", "type": "pic", "idx": 16, "box": (RIGHT - ZONE_W, M, ZONE_W, ZONE_H),
             "fill": None, "prompt": "PyCon NL logo (file and permission pending)", "sample": None},
            {"role": "qr", "name": "QR code", "type": "pic", "idx": 17,
             "box": (RIGHT - QR_TILE + QR_PAD, E(3.10) + QR_PAD, QR_TILE - 2 * QR_PAD, QR_TILE - 2 * QR_PAD),
             "fill": None, "prompt": "QR code", "sample": None},
        ],
    },
]

for _lay in LAYOUTS:
    if _lay["footer"]:
        _lay["ph"] += footer_phs(_lay["dark"])

# ------------------------------------------------------------------ parts


def theme_xml() -> str:
    clr = "".join(f'<a:{slot}><a:srgbClr val="{hexv}"/></a:{slot}>' for slot, hexv in CLR_SCHEME)

    def fnt(tag):
        return (f'<a:{tag}><a:latin typeface="{SANS}" panose="{PANOSE[SANS]}"/><a:ea typeface=""/>'
                f'<a:cs typeface=""/></a:{tag}>')

    solid = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    ln = "".join(f'<a:ln w="{w}" cap="flat" cmpd="sng" algn="ctr">{solid}<a:prstDash val="solid"/>'
                 f'<a:miter lim="800000"/></a:ln>' for w in (9525, 12700, 19050))
    eff = '<a:effectStyle><a:effectLst/></a:effectStyle>' * 3
    # objectDefaults: new shapes drawn in PowerPoint are flat blue with white Fira Sans; new lines are slate;
    # new text boxes use ink Fira Sans 20 pt (no colour or font outside the palette appears by default)
    sp_def = (
        '<a:spDef><a:spPr><a:solidFill><a:schemeClr val="accent1"/></a:solidFill><a:ln><a:noFill/></a:ln></a:spPr>'
        '<a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle><a:defPPr algn="ctr"><a:defRPr lang="en-GB" sz="2000">'
        '<a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:defRPr></a:defPPr></a:lstStyle>'
        '<a:style><a:lnRef idx="0"><a:schemeClr val="accent1"/></a:lnRef><a:fillRef idx="1"><a:schemeClr val="accent1"/>'
        '</a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef><a:fontRef idx="minor">'
        '<a:schemeClr val="lt1"/></a:fontRef></a:style></a:spDef>')
    ln_def = (
        '<a:lnDef><a:spPr><a:ln w="19050"><a:solidFill><a:schemeClr val="accent4"/></a:solidFill></a:ln></a:spPr>'
        '<a:bodyPr/><a:lstStyle/><a:style><a:lnRef idx="1"><a:schemeClr val="accent4"/></a:lnRef>'
        '<a:fillRef idx="0"><a:schemeClr val="accent4"/></a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent4"/>'
        '</a:effectRef><a:fontRef idx="minor"><a:schemeClr val="tx1"/></a:fontRef></a:style></a:lnDef>')
    tx_def = (
        '<a:txDef><a:spPr><a:noFill/></a:spPr><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0">'
        '<a:spAutoFit/></a:bodyPr><a:lstStyle><a:defPPr><a:defRPr lang="en-GB" sz="2000" dirty="0">'
        '<a:solidFill><a:schemeClr val="tx1"/></a:solidFill></a:defRPr></a:defPPr></a:lstStyle></a:txDef>')
    return (DECL + f'<a:theme xmlns:a="{NS["a"]}" name="{THEME_NAME}"><a:themeElements>'
            f'<a:clrScheme name="{THEME_NAME}">{clr}</a:clrScheme>'
            f'<a:fontScheme name="{THEME_NAME}">{fnt("majorFont")}{fnt("minorFont")}</a:fontScheme>'
            f'<a:fmtScheme name="{THEME_NAME}"><a:fillStyleLst>{solid * 3}</a:fillStyleLst>'
            f'<a:lnStyleLst>{ln}</a:lnStyleLst><a:effectStyleLst>{eff}</a:effectStyleLst>'
            f'<a:bgFillStyleLst>{solid * 3}</a:bgFillStyleLst></a:fmtScheme></a:themeElements>'
            f'<a:objectDefaults>{sp_def}{ln_def}{tx_def}</a:objectDefaults><a:extraClrSchemeLst/></a:theme>')


def level_styles(base_sz: float, colour: str, *, face: str, bullets: bool, aft: int, ln: int = 100) -> str:
    out = []
    for lvl in range(1, 10):
        if lvl == 1 or not bullets:
            out.append(ppr(lvl, marL=0, ln=ln, aft=aft,
                           rp=rpr("a:defRPr", sz=base_sz, colour=colour, face=face, kern=1200)))
        else:
            mar = 285750 * (lvl - 1)
            sz = 20
            out.append(ppr(lvl, marL=mar, indent=-285750, ln=ln, aft=aft,
                           bullet=("char", "•", "accent1" if lvl == 2 else "accent5"),
                           rp=rpr("a:defRPr", sz=sz, colour=colour, face=face, kern=1200)))
    return "".join(out)


def master_xml() -> str:
    shapes = [
        sp_placeholder(2, {"name": "Title", "type": "title", "idx": None, "box": (M, TITLE_Y, CW, TITLE_H1),
                           "anchor": "t", "prompt": "Title in sentence case, max 8 words"}, layout=True),
        sp_placeholder(3, {"name": "Body", "type": "body", "idx": 1, "box": (M, BODY_Y, CW, BOTTOM - BODY_Y),
                           "anchor": "t", "prompt": "Body text at 24 pt",
                           "levels": [("Body text at 24 pt", 0), ("Second level at 20 pt", 1),
                                      ("Third level at 20 pt", 2)]}, layout=True),
    ] + [sp_placeholder(4 + i, ph, layout=True) for i, ph in enumerate(footer_phs(False))]
    clrmap = ('<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
              'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>')
    ids = "".join(f'<p:sldLayoutId id="{2147483649 + i}" r:id="rId{i + 1}"/>' for i in range(len(LAYOUTS)))
    title_style = ("<p:titleStyle>" + ppr(1, ln=100, rp=rpr("a:defRPr", sz=32, b=False, colour="tx2", face=SANS_TITLE,
                                                             kern=1200)) + "</p:titleStyle>")
    body_style = "<p:bodyStyle>" + level_styles(24, "tx1", face="+mn-lt", bullets=True, aft=6) + "</p:bodyStyle>"
    other_style = ("<p:otherStyle><a:defPPr>" + rpr("a:defRPr") + "</a:defPPr>"
                   + level_styles(20, "tx1", face="+mn-lt", bullets=False, aft=0) + "</p:otherStyle>")
    return (DECL + f'<p:sldMaster {XMLNS}><p:cSld><p:bg><p:bgPr>{fill("bg1")}<a:effectLst/></p:bgPr></p:bg>'
            f'<p:spTree>{SPTREE_HEAD}{"".join(shapes)}</p:spTree></p:cSld>{clrmap}'
            f'<p:sldLayoutIdLst>{ids}</p:sldLayoutIdLst><p:hf hdr="0" dt="0"/>'
            f'<p:txStyles>{title_style}{body_style}{other_style}</p:txStyles></p:sldMaster>')


def sp_placeholder(sid: int, ph: dict, *, layout: bool) -> str:
    """A placeholder on the master or a layout (full geometry and list style)."""
    x, y, w, h = ph["box"]
    nv = (f'<p:nvSpPr><p:cNvPr id="{sid}" name="{escape(ph["name"])}"/><p:cNvSpPr><a:spLocks noGrp="1"/>'
          f'</p:cNvSpPr><p:nvPr>{ph_attr(ph.get("type"), ph.get("idx"), ph.get("sz"))}</p:nvPr></p:nvSpPr>')
    if ph.get("type") == "pic":
        sppr = (f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                + (fill(ph["fill"]) if ph.get("fill") else "<a:noFill/>") + '</p:spPr>')
        lst = ppr(1, algn="ctr", rp=rpr("a:defRPr", sz=14, colour="accent5"))
        body = (f'<p:txBody>{body_pr("ctr")}<a:lstStyle>{lst}</a:lstStyle>'
                f'{para(ph["prompt"])}</p:txBody>')
        return f"<p:sp>{nv}{sppr}{body}</p:sp>"
    sppr = f"<p:spPr>{xfrm(x, y, w, h)}</p:spPr>"
    lvl1 = ph.get("lvl1")
    lst = ""
    if ph.get("list") == "num":
        lst = ppr(1, marL=457200, indent=-457200, aft=10, bullet=("num", "arabicPeriod", "accent1"),
                  rp=rpr("a:defRPr", sz=24))
    elif lvl1 or ph.get("algn") or "aft" in ph:
        lst = ppr(1, algn=ph.get("algn", "l"), aft=ph.get("aft", 6 if ph.get("type") not in ("ftr", "sldNum") else 0),
                  rp=rpr("a:defRPr", **(lvl1 or {})))
    if ph.get("type") == "sldNum":
        content = (f'<a:p><a:fld id="{SLIDENUM_FLD}" type="slidenum">{rpr()}<a:t>‹#›</a:t></a:fld>'
                   f'<a:endParaRPr lang="{LANG}" dirty="0"/></a:p>')
    elif ph.get("levels"):
        content = "".join(para(t, lvl=lv) for t, lv in ph["levels"])
    else:
        content = para(ph["prompt"])
    return (f'<p:sp>{nv}{sppr}<p:txBody>{body_pr(ph.get("anchor", "t"))}<a:lstStyle>{lst}</a:lstStyle>'
            f'{content}</p:txBody></p:sp>')


def sp_static(sid: int, st: dict) -> str:
    x, y, w, h = st["box"]
    geom = (f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {st["radius"]}"/></a:avLst></a:prstGeom>'
            if st.get("radius") else '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>')
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(st["name"])}">{DECORATIVE}</p:cNvPr>'
            f'<p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
            f'<p:spPr>{xfrm(x, y, w, h)}{geom}{fill(st["colour"])}<a:ln><a:noFill/></a:ln></p:spPr></p:sp>')


# Navy layouts (L1, L3, L6, L11): the master's mapping, except hyperlinks, which map to sage (accent2). The theme's
# hlink is blue, 2.1:1 on navy (docs/DESIGN.md section 3); PowerPoint 2016 and older ignore the run-level
# hlinkClr extension the builder sets on the licence links and fall back to this mapping.
DARK_CLR_MAP = ('<p:clrMapOvr><a:overrideClrMapping bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" '
                'accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" '
                'hlink="accent2" folHlink="accent2"/></p:clrMapOvr>')
MASTER_CLR_MAP = '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>'


def layout_xml(i: int, lay: dict) -> str:
    shapes, sid = [], 2
    for st in lay["static"]:
        shapes.append(sp_static(sid, st))
        sid += 1
    for ph in lay["ph"]:
        shapes.append(sp_placeholder(sid, ph, layout=True))
        sid += 1
    bg = f'<p:bg><p:bgPr>{fill("tx2")}<a:effectLst/></p:bgPr></p:bg>' if lay["dark"] else ""
    type_attr = f' type="{lay["type"]}"' if lay["type"] else ""
    hf = "" if lay["footer"] else '<p:hf sldNum="0" hdr="0" ftr="0" dt="0"/>'
    if lay["footer"]:
        hf = '<p:hf hdr="0" dt="0"/>'
    return (DECL + f'<p:sldLayout {XMLNS} preserve="1" userDrawn="1"{type_attr}>'
            f'<p:cSld name="{escape(lay["name"])}">{bg}<p:spTree>{SPTREE_HEAD}{"".join(shapes)}</p:spTree></p:cSld>'
            f'{DARK_CLR_MAP if lay["dark"] else MASTER_CLR_MAP}{hf}</p:sldLayout>')


def slide_xml(i: int, lay: dict) -> str:
    """Example slide: every layout placeholder (pictures left empty), sample text, footer and slide number."""
    shapes, sid = [], 2
    for ph in lay["ph"]:
        nv = (f'<p:nvSpPr><p:cNvPr id="{sid}" name="{escape(ph["name"])}"/><p:cNvSpPr><a:spLocks noGrp="1"/>'
              f'</p:cNvSpPr><p:nvPr>{ph_attr(ph.get("type"), ph.get("idx"), ph.get("sz"))}</p:nvPr></p:nvSpPr>')
        sid += 1
        if ph.get("type") == "pic":
            shapes.append(f"<p:sp>{nv}<p:spPr/></p:sp>")
            continue
        if ph.get("type") == "sldNum":
            content = (f'<a:p><a:fld id="{SLIDENUM_FLD}" type="slidenum">{rpr()}<a:t>{i}</a:t></a:fld>'
                       f'<a:endParaRPr lang="{LANG}" dirty="0"/></a:p>')
        else:
            paras = []
            for item in ph["sample"]:
                text, lvl = (item, 0) if isinstance(item, str) else item
                paras.append(para(text, lvl=lvl))
            if ph["role"] == "title" and len(ph["sample"]) > 1:   # one paragraph, line break
                runs = f'<a:br>{rpr()}</a:br>'.join(f'<a:r>{rpr()}<a:t>{escape(t)}</a:t></a:r>' for t in ph["sample"])
                paras = [f'<a:p>{runs}<a:endParaRPr lang="{LANG}" dirty="0"/></a:p>']
            content = "".join(paras)
        shapes.append(f'<p:sp>{nv}<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>{content}</p:txBody></p:sp>')
    return (DECL + f'<p:sld {XMLNS}><p:cSld><p:spTree>{SPTREE_HEAD}{"".join(shapes)}</p:spTree></p:cSld>'
            f'{DARK_CLR_MAP if lay["dark"] else MASTER_CLR_MAP}</p:sld>')


def notes_master_xml() -> str:
    def ph(sid, name, typ, idx, box, sz=None, body=""):
        x, y, w, h = box
        sz_a = f' sz="{sz}"' if sz else ""
        if typ == "sldImg":
            return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1" '
                    f'noRot="1" noChangeAspect="1"/></p:cNvSpPr><p:nvPr><p:ph type="sldImg" idx="{idx}"/></p:nvPr>'
                    f'</p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/>'
                    f'<a:ln w="12700"><a:solidFill><a:schemeClr val="accent5"/></a:solidFill></a:ln></p:spPr></p:sp>')
        return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
                f'<p:nvPr><p:ph type="{typ}"{sz_a} idx="{idx}"/></p:nvPr></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}'
                f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody>'
                f'<a:bodyPr vert="horz" lIns="91440" tIns="45720" rIns="91440" bIns="45720" rtlCol="0"/>'
                f'<a:lstStyle/>{body}</p:txBody></p:sp>')
    num = (f'<a:p><a:pPr algn="r"/><a:fld id="{SLIDENUM_FLD}" type="slidenum">{rpr()}<a:t>‹#›</a:t></a:fld>'
           f'<a:endParaRPr lang="{LANG}" dirty="0"/></a:p>')
    shapes = [
        ph(2, "Slide image", "sldImg", 2, (381000, 685800, 6096000, 3429000)),
        ph(3, "Notes", "body", 3, (685800, 4343400, 5486400, 4114800), sz="quarter",
           body=para("Speaker notes: talk track, time, traceability")),
        ph(4, "Slide number", "sldNum", 5, (3884613, 8685213, 2971800, 458787), sz="quarter", body=num),
    ]
    style = "<p:notesStyle>" + level_styles(12, "tx1", face="+mn-lt", bullets=False, aft=4, ln=110) + "</p:notesStyle>"
    clrmap = ('<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
              'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>')
    return (DECL + f'<p:notesMaster {XMLNS}><p:cSld><p:bg><p:bgPr>{fill("bg1")}<a:effectLst/></p:bgPr></p:bg>'
            f'<p:spTree>{SPTREE_HEAD}{"".join(shapes)}</p:spTree></p:cSld>{clrmap}'
            f'<p:hf hdr="0" dt="0" ftr="0"/>{style}</p:notesMaster>')


def presentation_xml(n_slides: int) -> str:
    n_lay = len(LAYOUTS)
    sld_ids = "".join(f'<p:sldId id="{256 + i}" r:id="rId{n_lay + 10 + i}"/>' for i in range(n_slides))
    dts = "<p:defaultTextStyle><a:defPPr>" + rpr("a:defRPr") + "</a:defPPr>" + \
          level_styles(20, "tx1", face="+mn-lt", bullets=False, aft=0) + "</p:defaultTextStyle>"
    sec_ids = "".join(f'<p14:sldId id="{256 + i}"/>' for i in range(n_slides))
    ext = ('<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
           '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           '<p14:section name="Layout examples" id="{5E0C2B57-7A4C-4C55-9E5B-2026100300A1}">'
           f'<p14:sldIdLst>{sec_ids}</p14:sldIdLst></p14:section></p14:sectionLst></p:ext></p:extLst>')
    return (DECL + f'<p:presentation {XMLNS} saveSubsetFonts="0" embedTrueTypeFonts="1">'
            '<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
            '<p:notesMasterIdLst><p:notesMasterId r:id="rId2"/></p:notesMasterIdLst>'
            f'<p:sldIdLst>{sld_ids}</p:sldIdLst>'
            f'<p:sldSz cx="{SLIDE_W}" cy="{SLIDE_H}"/><p:notesSz cx="6858000" cy="9144000"/>'
            f'{dts}{ext}</p:presentation>')


R_OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_PML = "application/vnd.openxmlformats-officedocument.presentationml"


def rels(items: list[tuple[str, str, str]]) -> str:
    body = "".join(f'<Relationship Id="{i}" Type="{t}" Target="{tg}"/>' for i, t, tg in items)
    return DECL + f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'


def build_parts() -> dict[str, str]:
    n = len(LAYOUTS)
    parts: dict[str, str] = {}
    overrides = [
        ("/ppt/presentation.xml", f"{CT_PML}.presentation.main+xml"),
        ("/ppt/slideMasters/slideMaster1.xml", f"{CT_PML}.slideMaster+xml"),
        ("/ppt/notesMasters/notesMaster1.xml", f"{CT_PML}.notesMaster+xml"),
        ("/ppt/theme/theme1.xml", "application/vnd.openxmlformats-officedocument.theme+xml"),
        ("/ppt/theme/theme2.xml", "application/vnd.openxmlformats-officedocument.theme+xml"),
        ("/ppt/presProps.xml", f"{CT_PML}.presProps+xml"),
        ("/ppt/viewProps.xml", f"{CT_PML}.viewProps+xml"),
        ("/ppt/tableStyles.xml", f"{CT_PML}.tableStyles+xml"),
        ("/docProps/core.xml", "application/vnd.openxmlformats-package.core-properties+xml"),
        ("/docProps/app.xml", "application/vnd.openxmlformats-officedocument.extended-properties+xml"),
    ]
    for i in range(1, n + 1):
        overrides.append((f"/ppt/slideLayouts/slideLayout{i}.xml", f"{CT_PML}.slideLayout+xml"))
        overrides.append((f"/ppt/slides/slide{i}.xml", f"{CT_PML}.slide+xml"))
    ov = "".join(f'<Override PartName="{p}" ContentType="{c}"/>' for p, c in overrides)
    parts["[Content_Types].xml"] = (
        DECL + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>' + ov + "</Types>")
    parts["_rels/.rels"] = rels([
        ("rId1", f"{R_OFFICE}/officeDocument", "ppt/presentation.xml"),
        ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
         "docProps/core.xml"),
        ("rId3", f"{R_OFFICE}/extended-properties", "docProps/app.xml"),
    ])
    parts["docProps/core.xml"] = (
        DECL + '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<dc:title>PyCon NL 2026 style master</dc:title>'
        '<dc:subject>3D Data Science with Python, PyCon NL 2026</dc:subject>'
        '<dc:creator>Sergio Peral</dc:creator><cp:lastModifiedBy>Sergio Peral</cp:lastModifiedBy>'
        '<dc:description>Slide master in the PyCon NL 2026 palette (pycon-nl.org) with Fira Sans and Fira Mono. '
        'Built by deck/template/make_master.py.</dc:description>'
        f'<dc:language>{LANG}</dc:language><cp:revision>1</cp:revision>'
        '<dcterms:created xsi:type="dcterms:W3CDTF">2026-10-03T00:00:00Z</dcterms:created>'
        '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-03T00:00:00Z</dcterms:modified>'
        '</cp:coreProperties>')
    titles = "".join(f"<vt:lpstr>{escape(lay['name'])}</vt:lpstr>" for lay in LAYOUTS)
    parts["docProps/app.xml"] = (
        DECL + '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<TotalTime>0</TotalTime><Words>0</Words><Application>Microsoft Office PowerPoint</Application>'
        f'<PresentationFormat>Widescreen</PresentationFormat><Paragraphs>0</Paragraphs><Slides>{n}</Slides>'
        '<Notes>0</Notes><HiddenSlides>0</HiddenSlides><MMClips>0</MMClips><ScaleCrop>false</ScaleCrop>'
        '<HeadingPairs><vt:vector size="6" baseType="variant">'
        '<vt:variant><vt:lpstr>Fonts Used</vt:lpstr></vt:variant><vt:variant><vt:i4>3</vt:i4></vt:variant>'
        '<vt:variant><vt:lpstr>Theme</vt:lpstr></vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant>'
        f'<vt:variant><vt:lpstr>Slide Titles</vt:lpstr></vt:variant><vt:variant><vt:i4>{n}</vt:i4></vt:variant>'
        f'</vt:vector></HeadingPairs><TitlesOfParts><vt:vector size="{4 + n}" baseType="lpstr">'
        f'<vt:lpstr>{SANS}</vt:lpstr><vt:lpstr>{SANS_TITLE}</vt:lpstr><vt:lpstr>{MONO}</vt:lpstr>'
        f'<vt:lpstr>{THEME_NAME}</vt:lpstr>{titles}</vt:vector></TitlesOfParts>'
        '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>'
        '<AppVersion>16.0000</AppVersion></Properties>')
    pres_rels = [("rId1", f"{R_OFFICE}/slideMaster", "slideMasters/slideMaster1.xml"),
                 ("rId2", f"{R_OFFICE}/notesMaster", "notesMasters/notesMaster1.xml"),
                 ("rId3", f"{R_OFFICE}/theme", "theme/theme1.xml"),
                 ("rId4", f"{R_OFFICE}/presProps", "presProps.xml"),
                 ("rId5", f"{R_OFFICE}/viewProps", "viewProps.xml"),
                 ("rId6", f"{R_OFFICE}/tableStyles", "tableStyles.xml")]
    pres_rels += [(f"rId{n + 10 + i}", f"{R_OFFICE}/slide", f"slides/slide{i + 1}.xml") for i in range(n)]
    parts["ppt/_rels/presentation.xml.rels"] = rels(pres_rels)
    parts["ppt/presentation.xml"] = presentation_xml(n)
    parts["ppt/presProps.xml"] = DECL + f'<p:presentationPr {XMLNS}/>'
    parts["ppt/viewProps.xml"] = (DECL + f'<p:viewPr {XMLNS}><p:normalViewPr><p:restoredLeft sz="15620"/>'
                                  '<p:restoredTop sz="94660"/></p:normalViewPr><p:gridSpacing cx="76200" cy="76200"/>'
                                  '</p:viewPr>')
    parts["ppt/tableStyles.xml"] = (DECL + f'<a:tblStyleLst xmlns:a="{NS["a"]}" '
                                    'def="{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"/>')
    parts["ppt/theme/theme1.xml"] = theme_xml()
    parts["ppt/theme/theme2.xml"] = theme_xml()
    parts["ppt/slideMasters/slideMaster1.xml"] = master_xml()
    parts["ppt/slideMasters/_rels/slideMaster1.xml.rels"] = rels(
        [(f"rId{i + 1}", f"{R_OFFICE}/slideLayout", f"../slideLayouts/slideLayout{i + 1}.xml") for i in range(n)]
        + [(f"rId{n + 1}", f"{R_OFFICE}/theme", "../theme/theme1.xml")])
    parts["ppt/notesMasters/notesMaster1.xml"] = notes_master_xml()
    parts["ppt/notesMasters/_rels/notesMaster1.xml.rels"] = rels(
        [("rId1", f"{R_OFFICE}/theme", "../theme/theme2.xml")])
    for i, lay in enumerate(LAYOUTS, 1):
        parts[f"ppt/slideLayouts/slideLayout{i}.xml"] = layout_xml(i, lay)
        parts[f"ppt/slideLayouts/_rels/slideLayout{i}.xml.rels"] = rels(
            [("rId1", f"{R_OFFICE}/slideMaster", "../slideMasters/slideMaster1.xml")])
        parts[f"ppt/slides/slide{i}.xml"] = slide_xml(i, lay)
        parts[f"ppt/slides/_rels/slide{i}.xml.rels"] = rels(
            [("rId1", f"{R_OFFICE}/slideLayout", f"../slideLayouts/slideLayout{i}.xml")])
    return parts


def write_package(parts: dict[str, str], out: Path):
    for name, xml in parts.items():          # every part must be well-formed XML
        etree.fromstring(xml.encode("utf-8"))
    names = sorted(parts, key=lambda p: (p != "[Content_Types].xml", p))
    tmp = out.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in names:
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, parts[name].encode("utf-8"))
    os.replace(tmp, out)


def inch(v: int) -> float:
    return round(v / EMU_PER_IN, 4)


def geometry() -> dict:
    out = {"slide_in": [inch(SLIDE_W), inch(SLIDE_H)], "margin_in": inch(M), "right_edge_in": inch(RIGHT),
           "content_bottom_in": inch(BOTTOM), "media_aspect": "1588:1792 (397:448)",
           "media_emu": [MEDIA_W, MEDIA_H], "layouts": {}}
    for i, lay in enumerate(LAYOUTS, 1):
        items = []
        for st in lay["static"]:
            items.append({"kind": st["kind"], "name": st["name"], "colour": st["colour"],
                          "emu": list(st["box"]), "in": [inch(v) for v in st["box"]]})
        for ph in lay["ph"]:
            items.append({"kind": "placeholder", "name": ph["name"], "role": ph["role"],
                          "type": ph.get("type") or "obj", "idx": 0 if ph.get("idx") is None else ph["idx"],
                          "emu": list(ph["box"]), "in": [inch(v) for v in ph["box"]],
                          "anchor": ph.get("anchor"), "style": ph.get("lvl1")})
        out["layouts"][str(i)] = {"name": lay["name"], "dark": lay["dark"], "footer": lay["footer"],
                                  "shapes": items}
    return out


def check(out: Path):
    from pptx import Presentation
    prs = Presentation(str(out))
    assert prs.slide_width == SLIDE_W and prs.slide_height == SLIDE_H
    assert len(prs.slide_layouts) == len(LAYOUTS) == len(prs.slides)
    for i, (lay, s) in enumerate(zip(prs.slide_layouts, prs.slides), 1):
        assert lay.name == LAYOUTS[i - 1]["name"], (lay.name, i)
        assert s.slide_layout == lay
        idxs = sorted(p.placeholder_format.idx for p in lay.placeholders)
        print(f"L{i:<2} {lay.name:28} placeholders idx {idxs}")
    media = next(p for p in prs.slide_layouts[6].placeholders if p.placeholder_format.idx == 18)
    print(f"L7 media {media.width} x {media.height} EMU, ratio {media.width / media.height:.6f} "
          f"(1588/1792 = {1588 / 1792:.6f}), right edge {inch(media.left + media.width)} in")


def main():
    parts = build_parts()
    write_package(parts, OUT)
    GEOM_OUT.write_text(json.dumps(geometry(), indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    check(OUT)
    print(f"wrote {OUT.relative_to(HERE.parent.parent)} and {GEOM_OUT.name}")


if __name__ == "__main__":
    main()
